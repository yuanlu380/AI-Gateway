"""Generate a single combined PowerPoint with 4 demo summary slides (no cover pages)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Brand colors ──
DARK_BLUE = "003366"
ACCENT_BLUE = "0078D4"
ACCENT_TEAL = "00B294"
WHITE = "FFFFFF"
DARK_TEXT = "333333"
GRAY_TEXT = "666666"
ORANGE = "FF8C00"
RED_ACCENT = "E81123"
GREEN = "107C10"
PURPLE = "7B2D8B"
GRAY = "555555"

OUTPUT_DIR = r"c:\git\AI-Gateway\slides"
os.makedirs(OUTPUT_DIR, exist_ok=True)


def hex_to_rgb(h):
    return RGBColor(int(h[:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def add_box(slide, l, t, w, h, color, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb(color)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.color.rgb = hex_to_rgb(WHITE)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    return shape


def add_table(slide, headers, rows, left, top, width, height):
    tbl = slide.shapes.add_table(len(rows)+1, len(headers), Inches(left), Inches(top), Inches(width), Inches(height)).table
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i); cell.text = h; cell.fill.solid(); cell.fill.fore_color.rgb = hex_to_rgb(DARK_BLUE)
        for pg in cell.text_frame.paragraphs:
            pg.font.size = Pt(9); pg.font.color.rgb = hex_to_rgb(WHITE); pg.font.bold = True; pg.alignment = PP_ALIGN.CENTER
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r+1, c); cell.text = val
            cell.fill.solid(); cell.fill.fore_color.rgb = hex_to_rgb("F0F4F8" if r % 2 == 0 else WHITE)
            for pg in cell.text_frame.paragraphs:
                pg.font.size = Pt(9); pg.font.color.rgb = hex_to_rgb(DARK_TEXT); pg.alignment = PP_ALIGN.CENTER
    return tbl


def add_bullets(slide, left, top, width, height, lines, font_size=11):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "\u2022 " + line; p.font.size = Pt(font_size); p.font.color.rgb = hex_to_rgb(DARK_TEXT); p.space_before = Pt(3)
    return txBox


def add_slide_header(slide, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(0.06))
    bar.fill.solid(); bar.fill.fore_color.rgb = hex_to_rgb(ACCENT_BLUE); bar.line.fill.background()
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.55))
    p = txBox.text_frame.paragraphs[0]
    p.text = title; p.font.size = Pt(22); p.font.color.rgb = hex_to_rgb(DARK_BLUE); p.font.bold = True


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background; fill = bg.fill; fill.solid(); fill.fore_color.rgb = hex_to_rgb(WHITE)
    return slide


def add_text_block(slide, left, top, width, height, title_text, items, title_size=12, item_size=10):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title_text; p.font.size = Pt(title_size); p.font.color.rgb = hex_to_rgb(DARK_BLUE); p.font.bold = True
    for item in items:
        sp = tf.add_paragraph(); sp.text = item; sp.font.size = Pt(item_size); sp.font.color.rgb = hex_to_rgb(DARK_TEXT); sp.space_before = Pt(2)


# ═══════════════════════════════════════════
# Build the combined deck
# ═══════════════════════════════════════════
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# ── Slide 1: FinOps Framework ──
s = new_slide(prs)
add_slide_header(s, "FinOps Framework \u2014 Cost Management & Automated Governance")
for l, t, w, h, c, txt in [
    (0.3, 1.0, 1.4, 0.6, ACCENT_BLUE, "Client Apps"),
    (2.0, 1.0, 1.6, 0.6, DARK_BLUE, "Azure APIM"),
    (3.9, 1.0, 1.6, 0.6, ACCENT_TEAL, "Azure OpenAI"),
    (2.0, 1.9, 1.6, 0.6, ORANGE, "App Insights"),
    (3.9, 1.9, 1.6, 0.6, PURPLE, "Log Analytics"),
    (2.0, 2.8, 1.6, 0.6, RED_ACCENT, "Query Alert"),
    (3.9, 2.8, 1.6, 0.6, GREEN, "Logic App"),
]:
    add_box(s, l, t, w, h, c, txt)
add_table(s, ["Product", "TPM", "Quota", "Cost"],
    [["Platinum","2,000","1M","$15"],["Gold","1,000","1M","$10"],["Silver","500","1M","$5"]],
    5.8, 1.0, 3.9, 1.6)
add_bullets(s, 0.5, 4.1, 9, 3.2, [
    "emit-token-metric policy \u2192 App Insights (dimension: Product ID)",
    "KQL cost: (PromptTokens \u00d7 InputPrice + CompletionTokens \u00d7 OutputPrice) / 1000",
    "Scheduled Alert \u2192 Logic App auto-disables subscription exceeding CostQuota",
    "Models: gpt-4.1-mini, gpt-4.1, DeepSeek-V3.2  |  Per-product rate limiting per subscription",
    "Zero code changes \u2014 all cost tracking & governance enforced at gateway layer",
])

# ── Slide 2: AI Agent Service V3 ──
s = new_slide(prs)
add_slide_header(s, "AI Agent Service \u2014 Multi-Agent Orchestration via APIM Model Gateway")
for l, t, w, h, c, txt in [
    (0.3, 1.0, 1.5, 0.6, ACCENT_BLUE, "AI Foundry Hub"),
    (2.1, 1.0, 1.6, 0.6, PURPLE, "Agent Service"),
    (4.0, 1.0, 1.8, 0.6, DARK_BLUE, "APIM Gateway"),
    (4.0, 1.9, 1.8, 0.6, ORANGE, "Tool APIs"),
    (0.3, 1.9, 1.5, 0.6, GRAY, "Bing Search"),
    (2.1, 1.9, 1.6, 0.6, GREEN, "Logic App"),
]:
    add_box(s, l, t, w, h, c, txt)
add_table(s, ["Agent", "Tool", "Backend"],
    [["Math Tutor","Code Interpreter","Sandbox"],["Bing Assistant","Bing Grounding","Bing API"],
     ["Weather","OpenAPI Tool","APIM mock"],["Orders","OpenAPI Tool","Logic App"]],
    6.1, 1.0, 3.6, 2.0)
add_bullets(s, 0.5, 3.5, 9, 3.7, [
    "V3 Pattern: ai-gateway Foundry connection routes all LLM inference through APIM",
    "Model path: ai-gateway/gpt-4.1-mini \u2014 agents never call Azure OpenAI directly",
    "APIM converts subscription key \u2192 managed identity Bearer token (RBAC)",
    "emit-token-metric: 3 dimensions (Subscription ID, Client IP, API ID)",
    "Tool APIs (Weather, Orders, Catalog) also routed through APIM  |  Zero agent code changes",
])

# ── Slide 3: Token Rate Limiting ──
s = new_slide(prs)
add_slide_header(s, "Token Rate Limiting \u2014 Per-Subscription Token Budgets")
for l, t, w, h, c, txt in [
    (0.3, 1.1, 1.6, 0.65, ACCENT_BLUE, "Client (Key A)"),
    (0.3, 2.1, 1.6, 0.65, PURPLE, "Client (Key B)"),
    (2.4, 1.3, 2.2, 1.1, DARK_BLUE, "APIM\nllm-token-limit"),
    (5.0, 1.5, 2.0, 0.8, ACCENT_TEAL, "Azure OpenAI\ngpt-4.1-mini"),
]:
    add_box(s, l, t, w, h, c, txt)
add_table(s, ["Feature", "APIM Limit", "Platform Limit"],
    [["Granularity","Per sub/key","Per deployment"],["Isolation","Per counter-key","Global"],
     ["Custom tiers","Yes (Products)","No"],["Pre-flight","Yes (estimate)","No"]],
    5.5, 1.0, 4.2, 2.0)
add_bullets(s, 0.5, 3.5, 9, 3.7, [
    "Policy: llm-token-limit \u2014 self-contained, no App Insights required",
    "counter-key: context.Subscription.Id \u2014 isolates each subscription\u2019s token budget",
    "tokens-per-minute: configurable per deployment (demo: 100 TPM/subscription)",
    "HTTP 429 + Retry-After when exceeded \u2014 request never hits backend, counter resets every 60s",
    "Combinable with emit-token-metric for cost visibility, or Products for tiered limits",
])

# ── Slide 4: Model Gateway ──
s = new_slide(prs)
add_slide_header(s, "AI Foundry Model Gateway \u2014 Centralized AI Access & Security")
for l, t, w, h, c, txt in [
    (0.3, 1.0, 1.5, 0.6, PURPLE, "AI Foundry\nAgents"),
    (2.2, 1.0, 2.0, 0.6, DARK_BLUE, "APIM\nModel Gateway"),
    (4.6, 0.85, 1.8, 0.5, ACCENT_TEAL, "models-foundry\n(wt: 1)"),
    (4.6, 1.5, 1.8, 0.5, ORANGE, "agents-foundry\n(wt: 0)"),
    (6.8, 1.0, 1.5, 0.6, GREEN, "gpt-4o-mini\ngpt-4.1-mini"),
]:
    add_box(s, l, t, w, h, c, txt)
add_text_block(s, 0.5, 2.1, 4.8, 1.5, "3-Hop Authentication", [
    "1. Agent \u2192 APIM via ai-gateway connection (ApiKey)",
    "2. APIM validates sub key \u2192 managed identity token",
    "3. APIM \u2192 Cognitive Services via Bearer (RBAC)",
])
add_text_block(s, 5.5, 2.1, 4.2, 1.5, "Enterprise Capabilities", [
    "\u2705 Managed identity (zero keys)",
    "\u2705 Monitoring (App Insights)",
    "\u2705 Rate limiting per subscription",
    "\u2705 Semantic caching",
    "\u2705 Circuit breaker & retry",
])
add_bullets(s, 0.5, 4.1, 9, 3.2, [
    "ai-gateway connection (ApiManagement, isSharedToAll) \u2014 all agents use APIM as inference endpoint",
    "Backend pool: weight-based routing across Foundry instances (active/standby failover)",
    "Models: gpt-4o-mini (cap 10) + gpt-4.1-mini (cap 12)  |  Agents ref: ai-gateway/gpt-4.1-mini",
    "IaC (Bicep): fully reproducible \u2014 add models/backends without client changes",
])

# ── Save ──
path = os.path.join(OUTPUT_DIR, "AI-Gateway-Demos.pptx")
prs.save(path)
print(f"Created: {path}")
